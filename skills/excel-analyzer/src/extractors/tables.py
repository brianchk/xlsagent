"""Structured table (ListObject) extractor."""

from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from ..models import TableInfo
from .base import BaseExtractor


class TableExtractor(BaseExtractor):
    """Extracts structured table definitions from all sheets."""

    name = "tables"

    def extract(self) -> list[TableInfo]:
        """Extract all structured table definitions.

        Returns:
            List of TableInfo objects
        """
        tables = []

        for sheet_name in self.workbook.sheetnames:
            sheet = self.workbook[sheet_name]
            if not isinstance(sheet, Worksheet):
                continue

            sheet_tables = self._extract_sheet_tables(sheet, sheet_name)
            tables.extend(sheet_tables)

        return tables

    def _extract_sheet_tables(self, sheet: Worksheet, sheet_name: str) -> list[TableInfo]:
        """Extract tables from a sheet."""
        tables = []

        try:
            for table_name, table in sheet.tables.items():
                info = self._create_table_info(table, sheet_name, table_name)
                if info:
                    tables.append(info)
        except Exception:
            pass

        return tables

    def _create_table_info(self, table, sheet_name: str, table_name: str) -> TableInfo | None:
        """Create TableInfo from a table object."""
        try:
            # Get basic properties
            ref = table.ref if hasattr(table, "ref") else ""
            display_name = table.displayName if hasattr(table, "displayName") else table_name

            # Get column names
            columns = []
            try:
                if hasattr(table, "tableColumns") and table.tableColumns:
                    # Try different iteration methods
                    if hasattr(table.tableColumns, "tableColumn"):
                        for col in table.tableColumns.tableColumn:
                            if hasattr(col, "name"):
                                columns.append(col.name)
                    elif hasattr(table.tableColumns, "__iter__"):
                        for col in table.tableColumns:
                            if hasattr(col, "name"):
                                columns.append(col.name)
            except Exception:
                pass

            # If no columns found, try to get from sheet header row
            if not columns and ref:
                try:
                    sheet = self.workbook[sheet_name]
                    from openpyxl.utils import range_boundaries
                    min_col, min_row, max_col, max_row = range_boundaries(ref)
                    for col in range(min_col, max_col + 1):
                        cell = sheet.cell(row=min_row, column=col)
                        if cell.value:
                            columns.append(str(cell.value))
                except Exception:
                    pass

            # Get table style
            style_name = None
            try:
                if hasattr(table, "tableStyleInfo") and table.tableStyleInfo:
                    style_name = table.tableStyleInfo.name
            except Exception:
                pass

            # Check for totals row
            has_totals = False
            try:
                has_totals = getattr(table, "totalsRowShown", False) or False
            except Exception:
                pass

            # Check for header row
            has_header = True
            try:
                header_count = getattr(table, "headerRowCount", 1)
                has_header = header_count is None or header_count > 0
            except Exception:
                pass

            return TableInfo(
                name=table_name,
                sheet=sheet_name,
                range=ref,
                display_name=display_name,
                columns=columns,
                has_totals_row=has_totals,
                has_header_row=has_header,
                style_name=style_name,
            )
        except Exception:
            return None
