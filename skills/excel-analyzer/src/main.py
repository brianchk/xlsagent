"""Excel Workbook Analyzer - Extracts and documents Excel workbook contents."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

from .models import (
    WorkbookAnalysis,
    ExtractionError,
    ExtractionWarning,
)
from .extractors import (
    SheetExtractor,
    FormulaExtractor,
    NamedRangeExtractor,
    ConditionalFormatExtractor,
    DataValidationExtractor,
    PivotTableExtractor,
    ChartExtractor,
    TableExtractor,
    FilterExtractor,
    CommentExtractor,
    HyperlinkExtractor,
    ProtectionExtractor,
    PrintSettingsExtractor,
    ErrorExtractor,
    VBAExtractor,
    PowerQueryExtractor,
    ControlExtractor,
    ConnectionExtractor,
    DAXDetector,
)
from .reports.html_builder import HTMLReportBuilder
from .reports.markdown_builder import MarkdownReportBuilder


def capture_screenshots(file_path: Path, sheets: list, output_dir: Path) -> list:
    """Capture screenshots via Desktop Excel."""
    from .screenshots.desktop_excel import DesktopExcelScreenshotter
    screenshotter = DesktopExcelScreenshotter(output_dir / "screenshots")
    return screenshotter.capture_all_sheets(file_path, sheets)


def run_extractor(extractor_class, workbook, file_path, name: str):
    """Run a single extractor and handle errors."""
    try:
        extractor = extractor_class(workbook, file_path)
        return extractor.extract(), None
    except Exception as e:
        return None, ExtractionError(extractor=name, message=str(e))


def analyze_workbook(file_path: Path) -> WorkbookAnalysis:
    """Analyze an Excel workbook and extract all components."""
    print(f"Analyzing: {file_path.name}", flush=True)

    # Initialize analysis
    analysis = WorkbookAnalysis(
        file_path=file_path,
        file_name=file_path.name,
        file_size=file_path.stat().st_size,
    )

    # Load workbook for most extractors
    wb = load_workbook(file_path, data_only=False)

    # Run extractors
    extractors = [
        (SheetExtractor, "sheets"),
        (FormulaExtractor, "formulas"),
        (NamedRangeExtractor, "named_ranges"),
        (ConditionalFormatExtractor, "conditional_formats"),
        (DataValidationExtractor, "data_validations"),
        (PivotTableExtractor, "pivot_tables"),
        (ChartExtractor, "charts"),
        (TableExtractor, "tables"),
        (FilterExtractor, "auto_filters"),
        (CommentExtractor, "comments"),
        (HyperlinkExtractor, "hyperlinks"),
        (ProtectionExtractor, "protection"),
        (PrintSettingsExtractor, "print_settings"),
        (ErrorExtractor, "error_cells"),
    ]

    for extractor_class, attr_name in extractors:
        print(f"  Extracting: {attr_name}...", flush=True)
        result, error = run_extractor(extractor_class, wb, file_path, attr_name)
        if error:
            analysis.errors.append(error)
        elif result is not None:
            setattr(analysis, attr_name, result)

    wb.close()

    # VBA extraction (separate workbook load)
    print("  Extracting: VBA modules...", flush=True)
    try:
        vba_extractor = VBAExtractor(None, file_path)
        analysis.vba_modules = vba_extractor.extract()
        if analysis.vba_modules:
            analysis.has_macros = True
    except Exception as e:
        analysis.errors.append(ExtractionError(extractor="vba", message=str(e)))

    # Power Query extraction
    print("  Extracting: Power Query...", flush=True)
    try:
        pq_extractor = PowerQueryExtractor(None, file_path)
        analysis.power_queries = pq_extractor.extract()
    except Exception as e:
        analysis.errors.append(ExtractionError(extractor="power_query", message=str(e)))

    # Form controls
    print("  Extracting: Form controls...", flush=True)
    try:
        ctrl_extractor = ControlExtractor(None, file_path)
        analysis.form_controls = ctrl_extractor.extract()
    except Exception as e:
        analysis.errors.append(ExtractionError(extractor="controls", message=str(e)))

    # Data connections
    print("  Extracting: Data connections...", flush=True)
    try:
        conn_extractor = ConnectionExtractor(None, file_path)
        analysis.data_connections = conn_extractor.extract()
    except Exception as e:
        analysis.errors.append(ExtractionError(extractor="connections", message=str(e)))

    # DAX/Power Pivot detection
    print("  Detecting: DAX/Power Pivot...", flush=True)
    try:
        dax_detector = DAXDetector(None, file_path)
        analysis.has_dax, analysis.dax_detection_note = dax_detector.detect()
    except Exception as e:
        analysis.errors.append(ExtractionError(extractor="dax", message=str(e)))

    # Extract external references from formulas
    analysis.external_refs = list({
        ref
        for formula in analysis.formulas
        if formula.external_refs
        for ref in formula.external_refs
    })

    return analysis


def generate_reports(analysis: WorkbookAnalysis, output_dir: Path) -> None:
    """Generate HTML and Markdown reports."""
    print(f"Generating reports in: {output_dir}", flush=True)

    # HTML Report
    print("  Building HTML report...", flush=True)
    html_builder = HTMLReportBuilder(analysis)
    html_path = html_builder.build(output_dir)
    print(f"    Created: {html_path}", flush=True)

    # Markdown Reports
    print("  Building Markdown reports...", flush=True)
    md_builder = MarkdownReportBuilder(analysis)
    md_builder.build(output_dir)
    print(f"    Created: {output_dir}/README.md", flush=True)


def print_summary(analysis: WorkbookAnalysis) -> None:
    """Print analysis summary to console."""
    print("\n" + "=" * 60, flush=True)
    print(f"ANALYSIS COMPLETE: {analysis.file_name}", flush=True)
    print("=" * 60, flush=True)

    print(f"\nSheets: {len(analysis.sheets)}", flush=True)
    for s in analysis.sheets[:10]:
        vis = "" if s.visibility.value == "visible" else f" ({s.visibility.value})"
        print(f"  - {s.name}{vis}", flush=True)
    if len(analysis.sheets) > 10:
        print(f"  ... and {len(analysis.sheets) - 10} more", flush=True)

    print(f"\nFormulas: {len(analysis.formulas)}", flush=True)
    print(f"Named Ranges: {len(analysis.named_ranges)}", flush=True)
    print(f"  LAMBDA functions: {sum(1 for n in analysis.named_ranges if n.is_lambda)}", flush=True)

    if analysis.tables:
        print(f"Tables: {len(analysis.tables)}", flush=True)
    if analysis.pivot_tables:
        print(f"Pivot Tables: {len(analysis.pivot_tables)}", flush=True)
    if analysis.charts:
        print(f"Charts: {len(analysis.charts)}", flush=True)
    if analysis.vba_modules:
        print(f"VBA Modules: {len(analysis.vba_modules)}", flush=True)
    if analysis.power_queries:
        print(f"Power Queries: {len(analysis.power_queries)}", flush=True)
    if analysis.has_dax:
        print(f"DAX/Power Pivot: Detected ({analysis.dax_detection_note})", flush=True)

    if analysis.error_cells:
        print(f"\nError Cells: {len(analysis.error_cells)}", flush=True)
    if analysis.external_refs:
        print(f"External References: {len(analysis.external_refs)}", flush=True)

    if analysis.errors:
        print(f"\nExtraction Errors: {len(analysis.errors)}", flush=True)
        for e in analysis.errors[:5]:
            print(f"  - {e.extractor}: {e.message}", flush=True)

    if analysis.screenshots:
        print(f"\nScreenshots: {len(analysis.screenshots)}", flush=True)

    print(flush=True)


def main() -> int:
    """Main entry point."""
    parser = argparse.ArgumentParser(
        description="Analyze Excel workbooks and generate documentation"
    )
    parser.add_argument(
        "input",
        help="Path to Excel file (.xlsx, .xlsm, .xlsb)"
    )
    parser.add_argument(
        "-o", "--output",
        help="Output directory for reports (default: <filename>_analysis/)"
    )
    parser.add_argument(
        "--no-screenshots",
        action="store_true",
        help="Skip screenshot capture"
    )

    args = parser.parse_args()

    # Validate input file
    file_path = Path(args.input)
    if not file_path.exists():
        print(f"Error: File not found: {file_path}")
        return 1

    if not file_path.suffix.lower() in ('.xlsx', '.xlsm', '.xlsb'):
        print(f"Error: Not an Excel file: {file_path}")
        return 1

    # Determine output directory
    output_dir = Path(args.output) if args.output else Path.cwd() / f"{file_path.stem}_analysis"
    output_dir.mkdir(parents=True, exist_ok=True)

    try:
        # Run analysis
        analysis = analyze_workbook(file_path)

        # Capture screenshots
        if not args.no_screenshots:
            print("Capturing screenshots via Desktop Excel...", flush=True)
            try:
                screenshots = capture_screenshots(file_path, analysis.sheets, output_dir)
                if screenshots:
                    analysis.screenshots = screenshots
                    print(f"  Captured {len(screenshots)} screenshots", flush=True)
                else:
                    print("  No screenshots captured", flush=True)
            except Exception as e:
                print(f"  Warning: Could not capture screenshots: {e}", flush=True)
                analysis.warnings.append(ExtractionWarning(
                    extractor="screenshots",
                    message=f"Could not capture screenshots: {e}",
                ))

        # Generate reports
        generate_reports(analysis, output_dir)

        # Print summary
        print_summary(analysis)

        print(f"Reports saved to: {output_dir}", flush=True)
        print(f"  - HTML: {output_dir}/report.html", flush=True)
        print(f"  - Markdown: {output_dir}/README.md", flush=True)

        return 0

    except KeyboardInterrupt:
        print("\nCancelled.", flush=True)
        return 130
    except Exception as e:
        print(f"Error: {e}", flush=True)
        return 1


if __name__ == "__main__":
    sys.exit(main())
